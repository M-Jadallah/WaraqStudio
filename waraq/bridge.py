from __future__ import annotations

import json
import os
import re
import threading
import time
from pathlib import Path

from PySide6.QtCore import QObject, Property, QTimer, QUrl, Signal, Slot, Qt
from PySide6.QtGui import QDesktopServices, QGuiApplication
from PySide6.QtWidgets import QFileDialog

from . import __version__
from . import app_update
from .catalog import deepgram_models, opencode_models
from .config import (BUILTIN_PROVIDERS, credential_name, defaults, key_env_for, provider_catalog,
                     provider_info, provider_of)
from .models import DictListModel
from .remote import RemoteAccess
from .remote_commands import RemoteCommands
from .internet import InternetAccess
from .services import opencode_version, resolve_opencode
from .store import STAGES

LABELS = {"pending": "في الانتظار", "running": "قيد التنفيذ", "completed": "مكتمل", "skipped": "متجاوز",
          "failed": "فشل", "blocked": "يتطلب إجراء", "uncertain": "نتيجة غير مؤكدة", "paused": "متوقف", "cancelled": "ملغى"}
STAGE_LABELS = {"download": "التنزيل", "deepgram": "التفريغ", "opencode": "ملف Word"}


class Bridge(QObject):
    changed = Signal()
    jobsChanged = Signal()
    sourceChanged = Signal()
    settingsChanged = Signal()
    notice = Signal(str, bool)
    sourceReady = Signal(str)
    backgroundError = Signal(str)
    autoStopReached = Signal()
    libraryChanged = Signal()
    foldersChanged = Signal()
    youtubeVersionChanged = Signal()
    modelCatalogChanged = Signal()
    _modelCatalogReady = Signal(object, object, str, str)
    appUpdateChanged = Signal()
    _appUpdateReady = Signal(str, object)
    _appUpdateProgress = Signal(int)
    remoteAccessChanged = Signal()
    internetAccessChanged = Signal()
    accountsChanged = Signal()
    _accountTaskReady = Signal(str, int, str, object, str, int)
    _accountUsageReady = Signal(object)

    def __init__(self, store, config, secrets, scheduler):
        super().__init__()
        self.store, self.config, self.secrets, self.scheduler = store, config, secrets, scheduler
        self._jobs, self._workers, self._logs, self._source = [], [], [], []
        self._jobs_json = "[]"
        self._running = False
        self._source_title = ""
        self._busy = False
        self._account_ops = {}
        self._account_generations = {}
        self._account_pending = set()
        self._account_pending_owners = {}
        self._account_serial = 0
        self._account_logins = {}
        self._account_errors = {}
        self._accounts_closing = False
        self._accountTaskReady.connect(self._apply_account_task, Qt.ConnectionType.QueuedConnection)
        self._accountUsageReady.connect(self._apply_account_usage, Qt.ConnectionType.QueuedConnection)
        self._auto_stop = False
        self._auto_stop_seen = False
        self._auto_stop_fired = False
        self._library = []
        self._library_busy = False
        self._youtube_updating = False
        self._deepgram_models = []
        self._opencode_models = []
        self._catalog_errors = {"deepgram": "", "opencode": ""}
        self._catalog_busy = False
        self._update_release = None
        self._update_installer = None
        self._update_cancel = threading.Event()
        self._app_update = {"status": "idle", "version": "", "progress": 0,
                            "message": "افحص توفر إصدار جديد لتنزيله من GitHub.", "hasToken": False}
        try:
            self._app_update["hasToken"] = bool(secrets and secrets.get(app_update.TOKEN_KEY))
        except Exception:
            pass
        self._appUpdateReady.connect(self._apply_app_update)
        self._appUpdateProgress.connect(self._apply_app_update_progress)
        self.remote = RemoteAccess(self)
        self.remote.changed.connect(self.remoteAccessChanged)
        self.internet_commands = RemoteCommands(self)
        self.internet = InternetAccess(self, self.internet_commands.execute)
        self.internet.changed.connect(self.internetAccessChanged)
        self._internet_qr_url, self._internet_qr = "", ""
        self._modelCatalogReady.connect(self._apply_model_catalogs)
        # QML ListViews bind to these instead of the QVariantList properties:
        # rows are updated in place (dataChanged) so scroll position survives
        # every refresh tick, unlike a republished JS array which resets the view.
        self._jobs_model = DictListModel(parent=self)
        self._workers_model = DictListModel(parent=self)
        self._source_model = DictListModel(parent=self)
        self._results_model = DictListModel(parent=self)
        self._files_model = DictListModel(parent=self)
        self._providers_model = DictListModel(parent=self)
        self._accounts_model = DictListModel(parent=self)
        self._results_query = ""
        self._files_query = ""
        self._files_kind = "all"
        self._results_folder = "all"
        self._files_folder = "all"
        # refreshLibrary() may run on a worker thread (after organizing); the
        # queued connection applies the scan result to the model on the GUI thread.
        self.libraryChanged.connect(self._publish_files)
        self.settingsChanged.connect(self._publish_providers)
        self._publish_providers()
        self.discovery_cancel = threading.Event()
        self.sourceReady.connect(self._accept_source)
        self.backgroundError.connect(self._error)
        self.timer = QTimer(self)
        self.timer.timeout.connect(self.refresh)
        self.timer.start(700)
        # Only this independent timer schedules quota reads. Normal refresh and
        # phone snapshots remain cache-only, and quota work never reserves auth UI.
        self._account_usage_timer = QTimer(self)
        self._account_usage_timer.timeout.connect(self._schedule_account_usage)
        usage = getattr(self.scheduler, "account_usage", None)
        if usage:
            usage.set_callback(self._accountUsageReady.emit)
            self._account_usage_timer.start(1000)
        self.refresh()
        if self.config.migration_note:
            note = self.config.migration_note
            self.store.event("", note)
            # Delayed so the notice lands after the window is on screen.
            QTimer.singleShot(2500, lambda: self.notice.emit(note, False))

    @Property("QVariantList", notify=changed)
    def jobs(self): return self._jobs

    @Property(str, notify=jobsChanged)
    def jobsJson(self):
        """A cached snapshot QML can parse once into a native JavaScript array."""
        return self._jobs_json

    @Property("QVariantList", notify=changed)
    def workers(self): return self._workers

    @Property("QVariantList", notify=modelCatalogChanged)
    def deepgramModels(self): return self._deepgram_models

    @Property("QVariantList", notify=modelCatalogChanged)
    def opencodeModels(self): return self._opencode_models

    @Property("QVariantMap", notify=modelCatalogChanged)
    def catalogErrors(self): return self._catalog_errors

    @Property(bool, notify=modelCatalogChanged)
    def catalogBusy(self): return self._catalog_busy

    @Slot()
    def refreshModelCatalogs(self):
        if self._catalog_busy:
            return
        self._catalog_busy = True
        self.modelCatalogChanged.emit()
        cfg = self.config.snapshot()
        threading.Thread(target=self._load_model_catalogs, args=(cfg,), daemon=True).start()

    def _load_model_catalogs(self, cfg):
        dg, oc, dg_error, oc_error = None, None, "", ""
        try:
            dg = deepgram_models()
        except Exception:
            dg_error = "تعذّر جلب نماذج Deepgram. يمكنك إبقاء النموذج المحفوظ أو إعادة المحاولة."
        try:
            oc = opencode_models(cfg["opencode_executable"], cfg.get("opencode_context", ""))
        except Exception:
            oc_error = "تعذّر قراءة نماذج OpenCode المحلية. تحقق من مسار البرنامج ثم أعد المحاولة."
        self._modelCatalogReady.emit(dg, oc, dg_error, oc_error)

    @Slot(object, object, str, str)
    def _apply_model_catalogs(self, dg, oc, dg_error, oc_error):
        if dg is not None:
            self._deepgram_models = dg
        if oc is not None:
            self._opencode_models = oc
        self._catalog_errors = {"deepgram": dg_error, "opencode": oc_error}
        self._catalog_busy = False
        self.modelCatalogChanged.emit()

    @Property("QVariantList", notify=changed)
    def logs(self): return self._logs

    @Property(QObject, constant=True)
    def jobsModel(self): return self._jobs_model

    @Property(QObject, constant=True)
    def workersModel(self): return self._workers_model

    @Property(QObject, constant=True)
    def sourceModel(self): return self._source_model

    @Property(QObject, constant=True)
    def resultsModel(self): return self._results_model

    @Property(QObject, constant=True)
    def filesModel(self): return self._files_model

    @Property(QObject, constant=True)
    def providersModel(self): return self._providers_model

    @Property(QObject, constant=True)
    def accountsModel(self): return self._accounts_model

    @Property(bool, notify=accountsChanged)
    def accountBusy(self):
        return bool(self._account_ops or self._account_pending or self._account_logins)

    def _usage_snapshot(self, ident):
        service = getattr(self.scheduler, "account_usage", None)
        if service:
            return service.public_snapshot(ident)
        window = {"status": "unknown", "remainingPercent": None, "resetsAt": None, "windowMinutes": None}
        return {"status": "unknown", "stale": False, "lastSuccessAt": None, "lastAttemptAt": None,
                "messageCode": "not_fetched", "source": "", "refreshing": False,
                "fiveHour": dict(window), "weekly": dict(window)}

    @Slot(str)
    def refreshAccountUsage(self, ident):
        service = getattr(self.scheduler, "account_usage", None)
        if service and not self._accounts_closing and self._account(ident):
            service.request(ident, manual=True)
            self._publish_accounts()

    @Slot()
    def refreshAllAccountUsage(self):
        service = getattr(self.scheduler, "account_usage", None)
        if service and not self._accounts_closing:
            service.request_all()
            self._publish_accounts()

    @Slot(object)
    def _apply_account_usage(self, payload):
        service = getattr(self.scheduler, "account_usage", None)
        if service and not self._accounts_closing and service.apply(payload):
            self._publish_accounts()

    @Slot()
    def _schedule_account_usage(self):
        service = getattr(self.scheduler, "account_usage", None)
        if service and not self._accounts_closing:
            service.schedule_due()
            self._publish_accounts()

    def _invalidate_account_usage(self, ident):
        service = getattr(self.scheduler, "account_usage", None)
        if service:
            service.invalidate(ident)

    def _publish_accounts(self):
        registry = getattr(self.scheduler, "accounts", None)
        rows = []
        states = {"authenticated": "مسجّل الدخول", "complete": "مسجّل الدخول", "ready": "جاهز",
                  "authorizing": "بانتظار الموافقة في المتصفح", "pending": "بانتظار الموافقة في المتصفح",
                  "failed": "تعذّر تسجيل الدخول", "expired": "انتهت مهلة الدخول", "signed_out": "يلزم تسجيل الدخول",
                  "logged_out": "يلزم تسجيل الدخول", "duplicate": "الحساب مضاف مسبقًا", "needsverify": "يلزم التحقق من الهوية",
                  "unknown": "لم يُفحص", "cooldown": "بانتظار تجدد الحد", "removing": "جارٍ حذف الدخول",
                  "checking": "جارٍ فحص الحساب"}
        for account in registry.list() if registry else []:
            ident = account["id"]
            # Explicit projection: native credential responses never enter a QVariant/QML model.
            state = str(account.get("login_state") or "signed_out")
            health = str(account.get("health") or "")
            error = self._account_errors.get(ident, "") or str(account.get("error") or "")
            cooldown = int(account.get("cooldown_seconds") or 0)
            operation = self._account_ops.get(ident, "")
            login = self._account_logins.get(ident, {})
            usage = self._usage_snapshot(ident)
            service = getattr(self.scheduler, "account_usage", None)
            usage_maintenance = bool(service and service.state(ident).get("maintenance"))
            usage_refreshing = bool(usage.get("refreshing"))
            usage = {key: usage.get(key) for key in ("status", "stale", "lastSuccessAt", "lastAttemptAt",
                                                    "messageCode", "source", "fiveHour", "weekly")}
            for key in ("fiveHour", "weekly"):
                window = usage[key] or {}
                usage[key] = {field: window.get(field) for field in ("status", "remainingPercent", "resetsAt", "windowMinutes")}
            status_text = states.get(state, state)
            if state == "authenticated":
                status_text = {"model_limited": "بعض النماذج غير متاحة مؤقتًا",
                               "cooldown": "بانتظار تجدد حد الاستخدام",
                               "probe": "حان موعد إعادة فحص الحد؛ لم يُثبت تجدده"}.get(health, status_text)
            if operation == "metadata":
                status_text = "جارٍ تحديث الحساب…"
            if health == "execution_uncertain":
                status_text = "جلسة سابقة لم يُتحقق من توقفها"
            if operation == "recheck":
                status_text = "جارٍ التحقق من توقف الجلسة…"
            rows.append({"id": ident, "alias": str(account.get("alias") or "حساب ChatGPT"),
                         "priority": int(account.get("priority") or 0), "enabled": bool(account.get("enabled")),
                         "preferred": bool(account.get("preferred")), "email": str(account.get("email") or ""),
                         "uncertain": health == "execution_uncertain",
                         "subscription": str(account.get("subscription") or ""),
                         "inFlight": int(account.get("in_flight") or 0),
                         "busy": bool(operation or ident in self._account_pending_owners.values() or usage_maintenance),
                         "loggingIn": operation == "login", "needsCode": login.get("mode") == "code",
                         "statusText": status_text,
                         "usageRefreshing": usage_refreshing, "usageMaintenance": usage_maintenance, "usage": usage,
                         "cooldownText": ("إعادة فحص توفر التنفيذ بعد " + str((cooldown + 59) // 60) + " دقيقة") if cooldown else "",
                         "error": error, "quotaText": ("يُعاد الفحص بعد " + str((cooldown + 59) // 60) + " دقيقة")
                         if cooldown else "الحصة المتبقية غير معروفة"})
        account_changed = self._accounts_model._items != rows
        self._accounts_model.update_items(rows)
        if account_changed:
            self.accountsChanged.emit()

    def _account(self, ident):
        registry = getattr(self.scheduler, "accounts", None)
        return next((row for row in registry.list() if row["id"] == ident), None) if registry else None

    def _account_mutable(self, ident):
        account = self._account(ident)
        service = getattr(self.scheduler, "account_usage", None)
        maintenance = bool(service and service.state(ident).get("maintenance"))
        return bool(account and not account.get("in_flight") and ident not in self._account_ops
                    and ident not in self._account_pending_owners.values()
                    and not maintenance and not self._accounts_closing)

    @Slot(str)
    def addAccount(self, alias):
        try:
            self.scheduler.accounts.create(alias.strip())
            self._publish_accounts()
        except Exception:
            self.notice.emit("تعذّرت إضافة الحساب. اختر اسمًا من 1 إلى 80 محرفًا.", True)

    @Slot(str, str)
    def renameAccount(self, ident, alias):
        if not self._account_mutable(ident):
            return
        try:
            self.scheduler.accounts.update(ident, alias=alias.strip())
            self._publish_accounts()
        except Exception:
            self.notice.emit("تعذّر تغيير اسم الحساب.", True)

    @Slot(str, bool)
    def setAccountEnabled(self, ident, enabled):
        if ident in self._account_ops:
            return
        try:
            self.scheduler.accounts.update(ident, enabled=enabled)
            self._publish_accounts()
        except Exception:
            self.notice.emit("تعذّر تحديث تفعيل الحساب.", True)

    @Slot(str)
    def preferAccount(self, ident):
        try:
            self.scheduler.accounts.prefer(ident)
            self._publish_accounts()
        except Exception:
            self.notice.emit("تعذّر تفضيل الحساب.", True)

    @Slot(str, int)
    def moveAccount(self, ident, direction):
        try:
            rows = self.scheduler.accounts.list()
            index = next(i for i, item in enumerate(rows) if item["id"] == ident)
            target = index + (1 if direction > 0 else -1)
            if 0 <= target < len(rows):
                rows[index], rows[target] = rows[target], rows[index]
                with self.store.lock:
                    for priority, account in enumerate(rows, 1):
                        self.scheduler.accounts.update(account["id"], priority=priority)
                self._publish_accounts()
        except Exception:
            self.notice.emit("تعذّر تغيير ترتيب الحساب.", True)

    def _account_task(self, ident, generation, action, work):
        self._account_serial += 1
        token = self._account_serial
        self._account_pending.add(token)
        self._account_pending_owners[token] = ident
        self.accountsChanged.emit()
        def run():
            result, error = None, ""
            try:
                result = work()
            except Exception:
                # Native exceptions may contain authorization URLs, codes or credentials.
                error = {"begin": "تعذّر بدء الدخول. تحقّق من دعم تسجيل ChatGPT في OpenCode المثبّت.",
                         "finish": "تعذّر إكمال الدخول. أعد المحاولة ووافق في المتصفح.",
                         "metadata": "تعذّر فحص الحساب. تحقّق من اتصال OpenCode وأعد المحاولة.",
                         "recheck": "لم يُثبت توقف الجلسة السابقة. انتظر ثم أعد فحص الجلسة قبل إعادة التنفيذ.",
                         "remove": "تعذّر حذف تسجيل الدخول المحلي. أعد المحاولة.",
                         "cancel": "تعذّر إلغاء محاولة الدخول. أعد المحاولة."}.get(action, "تعذّر تنظيف محاولة الدخول.")
            self._accountTaskReady.emit(ident, generation, action, result, error, token)
        try:
            threading.Thread(target=run, daemon=True, name="account-" + action).start()
        except RuntimeError:
            self._accountTaskReady.emit(ident, generation, action, None, "تعذّر بدء العملية في الخلفية. أعد المحاولة.", token)

    @Slot(str)
    def loginAccount(self, ident):
        if not self._account_mutable(ident) or self.accountBusy:
            return
        try:
            self.scheduler.accounts.set_login_state(ident, "authorizing")
        except Exception:
            return
        self._invalidate_account_usage(ident)
        generation = self._account_generations.get(ident, 0) + 1
        self._account_generations[ident] = generation
        self._account_ops[ident] = "login"
        self._account_errors.pop(ident, None)
        self._publish_accounts()
        self._account_task(ident, generation, "begin", lambda: self.scheduler.account_runtime.begin_login(ident))

    @Slot(str)
    def recheckAccount(self, ident):
        if not self._account_mutable(ident):
            return
        generation = self._account_generations.get(ident, 0) + 1
        self._account_generations[ident] = generation
        self._account_ops[ident] = "recheck"
        self._account_task(ident, generation, "recheck", lambda: self.scheduler.recheck_account(ident))
        self._publish_accounts()

    @Slot(str, str)
    def finishAccountLogin(self, ident, code):
        login = self._account_logins.get(ident)
        if not login or login.get("polling") or self._accounts_closing:
            return
        if login.get("mode") == "code" and not code.strip():
            return
        login["polling"] = True
        generation = self._account_generations[ident]
        self._account_task(ident, generation, "finish",
                           lambda: self.scheduler.account_runtime.finish_login(ident, code=code.strip() or None))

    @Slot(str)
    def cancelAccountLogin(self, ident):
        if self._account_ops.get(ident) != "login":
            return
        generation = self._account_generations.get(ident, 0) + 1
        self._account_generations[ident] = generation
        self._account_logins.pop(ident, None)
        self._account_ops[ident] = "cancel"
        self._account_task(ident, generation, "cancel", lambda: self.scheduler.account_runtime.cancel_login(ident))
        self._publish_accounts()

    @Slot(str)
    def refreshAccount(self, ident):
        if not self._account_mutable(ident):
            return
        try:
            self.scheduler.accounts.set_login_state(ident, "checking")
        except Exception:
            return
        self._account_ops[ident] = "metadata"
        generation = self._account_generations.get(ident, 0) + 1
        self._account_generations[ident] = generation
        self._account_task(ident, generation, "metadata", lambda: self.scheduler.account_runtime.metadata(ident))
        self._publish_accounts()

    @Slot(str)
    def removeAccount(self, ident):
        if not self._account_mutable(ident):
            return
        try:
            self.scheduler.accounts.set_login_state(ident, "removing")
        except Exception:
            return
        self._invalidate_account_usage(ident)
        generation = self._account_generations.get(ident, 0) + 1
        self._account_generations[ident] = generation
        self._account_ops[ident] = "remove"
        self._account_task(ident, generation, "remove", lambda: self.scheduler.account_runtime.delete_auth(ident))
        self._publish_accounts()

    @staticmethod
    def _trusted_auth_url(value):
        url = QUrl(str(value))
        return (url.isValid() and url.scheme() == "https" and url.host().lower() == "auth.openai.com"
                and url.port(-1) in (-1, 443) and not url.userInfo())

    @Slot(str, int, str, object, str, int)
    def _apply_account_task(self, ident, generation, action, result, error, token):
        self._account_pending.discard(token)
        self._account_pending_owners.pop(token, None)
        if self._accounts_closing or self._account_generations.get(ident) != generation or not self._account(ident):
            self.accountsChanged.emit()
            return
        registry = self.scheduler.accounts
        if action == "cleanup":
            self._publish_accounts()
            self.accountsChanged.emit()
            return
        try:
            if error:
                raise ValueError(error)
            if action == "begin":
                if not isinstance(result, dict) or not self._trusted_auth_url(result.get("url", "")):
                    error = "رُفض عنوان تسجيل الدخول؛ يجب أن يكون اتصال HTTPS موثوقًا لدى OpenAI."
                    raise ValueError(error)
                if result.get("mode") not in ("auto", "code"):
                    error = "طريقة الدخول التي أعادها OpenCode غير مدعومة. حدّث البرنامج وأعد المحاولة."
                    raise ValueError(error)
                self._account_logins[ident] = {"mode": result["mode"], "polling": False}
                if not QDesktopServices.openUrl(QUrl(result["url"])):
                    error = "تعذّر فتح المتصفح لتسجيل الدخول. تحقّق من المتصفح الافتراضي وأعد المحاولة."
                    raise ValueError(error)
                if result["mode"] == "auto":
                    self.finishAccountLogin(ident, "")
            elif action == "finish" and isinstance(result, dict) and result.get("status") == "pending":
                self._account_logins[ident]["polling"] = False
                self._account_logins[ident]["mode"] = "auto"
                QTimer.singleShot(1800, lambda: self.finishAccountLogin(ident, "")
                                 if self._account_generations.get(ident) == generation else None)
            else:
                if action in ("finish", "metadata"):
                    if not isinstance(result, dict) or not result.get("authenticated"):
                        error = "لم يكتمل تسجيل الدخول. أعد المحاولة ووافق في المتصفح."
                        raise ValueError(error)
                    registry.record_login_result(ident, result, reset_cooldown=action != "metadata")
                elif action == "remove":
                    registry.remove(ident)
                elif action == "cancel":
                    registry.set_login_state(ident, "signed_out")
                self._account_ops.pop(ident, None)
                self._account_logins.pop(ident, None)
                self._account_errors.pop(ident, None)
        except Exception:
            self._account_ops.pop(ident, None)
            self._account_logins.pop(ident, None)
            self._account_errors[ident] = error or "لم تكتمل العملية. أعد المحاولة أو ألغِ تسجيل الدخول."
            account = self._account(ident)
            if account and account.get("login_state") == "duplicate":
                self._account_errors[ident] = "هذا الحساب مسجّل بالفعل باسم آخر. استخدم الإدخال الموجود."
            elif action != "recheck" and account and account.get("health") != "execution_uncertain":
                registry.set_login_state(ident, "failed", self._account_errors[ident])
            if action in ("begin", "finish"):
                self._account_task(ident, generation, "cleanup", lambda: self.scheduler.account_runtime.cancel_login(ident))
        self._publish_accounts()
        self.accountsChanged.emit()

    @Slot()
    def closeAccounts(self):
        # Invalidate callbacks before scheduler.shutdown closes native runtimes.
        self._accounts_closing = True
        self._account_usage_timer.stop()
        service = getattr(self.scheduler, "account_usage", None)
        if service:
            service.stop()
        self._account_generations = {key: value + 1 for key, value in self._account_generations.items()}

    @Property(bool, notify=sourceChanged)
    def hasSourceSelection(self): return any(row.get("selected") for row in self._source)

    @Property(int, notify=libraryChanged)
    def libraryCount(self): return len(self._library)

    @Property("QVariantList", notify=sourceChanged)
    def source(self): return self._source

    @Property(str, notify=sourceChanged)
    def sourceTitle(self): return self._source_title

    @Property(bool, notify=sourceChanged)
    def busy(self): return self._busy

    @Property(bool, notify=changed)
    def running(self): return self.scheduler.running

    @Property(str, notify=settingsChanged)
    def settingsJson(self): return json.dumps(self.config.snapshot(), ensure_ascii=False, indent=2)

    @Property(str, constant=True)
    def version(self): return __version__

    @Property("QVariantMap", notify=appUpdateChanged)
    def appUpdate(self): return dict(self._app_update)

    @Property("QVariantMap", notify=remoteAccessChanged)
    def remoteAccess(self): return self.remote.state()

    @Property("QVariantMap", notify=internetAccessChanged)
    def internetAccess(self): return self.internet.state()

    @Property(str, notify=internetAccessChanged)
    def internetQr(self):
        url = self.internet.state().get("pairUrl", "")
        if url != self._internet_qr_url:
            self._internet_qr_url, self._internet_qr = url, ""
            if url:
                import qrcode
                import qrcode.image.svg
                from urllib.parse import quote
                svg = qrcode.make(url, image_factory=qrcode.image.svg.SvgPathImage, border=4).to_string().decode("utf-8")
                self._internet_qr = "data:image/svg+xml," + quote(svg, safe="")
        return self._internet_qr

    @Slot(str, str, str)
    def registerInternet(self, url, code, name):
        self.internet.register(url, code, name)

    @Slot()
    def startInternet(self):
        self.internet.start()

    @Slot()
    def stopInternet(self):
        self.internet.stop()
        self.internet_commands.stop()

    @Slot()
    def closeRemoteLibrary(self):
        """Cancel shared library work only when the desktop itself exits."""
        self._remote_library.close()

    @Slot()
    def createInternetPair(self):
        self.internet.createPair()

    @Slot(str)
    def approveInternetPhone(self, claim_id):
        self.internet.approvePair(claim_id)

    @Slot(str)
    def rejectInternetPhone(self, claim_id):
        self.internet.rejectPair(claim_id)

    @Slot()
    def refreshInternetPhones(self):
        self.internet.refreshPhones()

    @Slot(str)
    def revokeInternetPhone(self, phone_id):
        self.internet.revokePhone(phone_id)

    @Slot()
    def forgetInternetDevice(self):
        self.internet.forgetDevice()

    @Slot(int)
    def startRemote(self, port):
        self.remote.start(port)

    @Slot()
    def stopRemote(self):
        self.remote.stop()

    @Slot()
    def renewRemoteCode(self):
        if self.remote.state()["enabled"]:
            self.remote.rotate_code()

    @Slot()
    def refreshRemoteAddresses(self):
        if self.remote.state()["enabled"]:
            self.remote.refresh_addresses()

    def _set_app_update(self, **values):
        self._app_update.update(values)
        self.appUpdateChanged.emit()

    @Slot(str)
    def saveUpdateToken(self, token):
        if self._app_update["status"] in ("checking", "downloading"):
            self.notice.emit("انتظر انتهاء عملية التحديث قبل تغيير رمز GitHub.", True)
            return
        try:
            if not self.secrets:
                raise RuntimeError("مخزن بيانات الاعتماد غير متاح.")
            token = str(token).strip()
            if token:
                self.secrets.set(app_update.TOKEN_KEY, token)
            else:
                self.secrets.delete(app_update.TOKEN_KEY)
            self._set_app_update(hasToken=bool(token))
            self.notice.emit("حُفظ رمز GitHub في مخزن بيانات الاعتماد." if token else "أُزيل رمز GitHub المحفوظ.", False)
        except Exception:
            self.notice.emit("تعذر تعديل رمز GitHub في مخزن بيانات الاعتماد.", True)

    @Slot()
    def checkAppUpdate(self):
        if self._app_update["status"] in ("checking", "downloading"):
            return
        self._update_release = self._update_installer = None
        self._set_app_update(status="checking", progress=0, version="", message="جارٍ فحص آخر إصدار منشور…")
        def work():
            try:
                token = self.secrets.get(app_update.TOKEN_KEY) if self.secrets else ""
                release = app_update.check_latest(self.version, token)
                self._appUpdateReady.emit("checked", release)
            except app_update.UpdateError as exc:
                self._appUpdateReady.emit("error", str(exc))
            except Exception:
                self._appUpdateReady.emit("error", "تعذر فحص التحديث أو قراءة رمز GitHub. تحقق من الاتصال ومخزن بيانات الاعتماد.")
        try:
            threading.Thread(target=work, daemon=True, name="app-update-check").start()
        except Exception:
            self._apply_app_update("error", "تعذر بدء فحص التحديث. أعد المحاولة.")

    @Slot()
    def downloadAppUpdate(self):
        if self._app_update["status"] != "available" or not self._update_release:
            return
        release = dict(self._update_release)
        self._update_cancel = threading.Event()
        cancel = self._update_cancel
        self._set_app_update(status="downloading", progress=0, message="جارٍ تنزيل مُثبّت الإصدار " + release["version"] + "…")
        def work():
            try:
                token = self.secrets.get(app_update.TOKEN_KEY) if self.secrets else ""
                path = app_update.download_installer(self.config.home, release, cancel,
                                                     self._appUpdateProgress.emit, token)
                self._appUpdateReady.emit("downloaded", str(path))
            except app_update.UpdateCancelled as exc:
                self._appUpdateReady.emit("cancelled", str(exc))
            except app_update.UpdateError as exc:
                self._appUpdateReady.emit("error", str(exc))
            except Exception:
                self._appUpdateReady.emit("error", "تعذر تنزيل التحديث. تحقق من الاتصال والمساحة ومخزن بيانات الاعتماد.")
        try:
            threading.Thread(target=work, daemon=True, name="app-update-download").start()
        except Exception:
            self._apply_app_update("error", "تعذر بدء تنزيل التحديث. أعد فحص الإصدار.")

    @Slot()
    def cancelAppUpdate(self):
        if self._app_update["status"] == "downloading":
            self._update_cancel.set()
            self._set_app_update(message="جارٍ إلغاء التنزيل…")

    @Slot(int)
    def _apply_app_update_progress(self, progress):
        if self._app_update["status"] == "downloading":
            self._set_app_update(progress=max(0, min(100, progress)))

    @Slot(str, object)
    def _apply_app_update(self, action, result):
        if action == "checked":
            self._update_release = result if result["available"] else None
            version = result["version"]
            self._set_app_update(status="available" if result["available"] else "up_to_date", version=version,
                                 message="يتوفر الإصدار " + version + " للتنزيل." if result["available"] else "نسختك الحالية هي الأحدث.")
        elif action == "downloaded":
            self._update_installer = Path(result)
            self._set_app_update(status="downloaded", progress=100,
                                 message="اكتمل التنزيل والتحقق. أوقف المعالجة ثم شغّل المُثبّت؛ سيُغلق التطبيق وتُحفظ بياناتك.")
        elif action == "cancelled":
            self._set_app_update(status="available", progress=0, message=str(result))
        else:
            self._set_app_update(status="error", message=str(result))

    @Slot()
    def installAppUpdate(self):
        if self._app_update["status"] != "downloaded" or not self._update_installer or not self._update_release:
            return
        # Pause prevents new dispatch; active captures publication after completion.
        if self.scheduler.running or not self.scheduler.idle() or self._busy or self._library_busy or self._youtube_updating or self.accountBusy:
            self.notice.emit("أوقف الطابور وانتظر انتهاء العمليات النشطة قبل تشغيل مُثبّت التحديث.", True)
            return
        if os.name != "nt":
            self.notice.emit("تشغيل مُثبّت التحديث متاح على Windows فقط.", True)
            return
        try:
            app_update.verify_installer(self._update_installer, self._update_release)
            if self.scheduler.running or not self.scheduler.idle() or self._busy or self._library_busy or self._youtube_updating or self.accountBusy:
                self.notice.emit("بدأت عملية معالجة أثناء التحقق؛ انتظر توقفها ثم أعد تثبيت التحديث.", True)
                return
            os.startfile(str(self._update_installer), "runas")
        except app_update.UpdateError as exc:
            self._set_app_update(status="error", message=str(exc))
            return
        except Exception:
            self.notice.emit("تعذر تشغيل المُثبّت أو أُلغي طلب صلاحيات Windows. يمكنك المحاولة مرة أخرى.", True)
            return
        self.discovery_cancel.set()
        self.scheduler.shutdown()
        QGuiApplication.quit()

    @Property(float, notify=settingsChanged)
    def uiScale(self): return float(self.config.snapshot().get("ui_scale", 1.0))

    @Property("QVariantList", notify=settingsChanged)
    def templates(self):
        return [{"id": t["id"], "name": t["name"], "text": t["text"]} for t in self.config.snapshot()["opencode_templates"]]

    @Property("QVariantList", notify=settingsChanged)
    def providers(self):
        """Mirror of providersModel for ComboBoxes and Python-side checks."""
        return self._provider_rows()

    @Property("QVariantList", notify=libraryChanged)
    def library(self): return self._library

    @Property("QVariantList", notify=foldersChanged)
    def folders(self):
        return [{"id": "", "name": "بدون مجلد"}] + self.store.folders()

    @Slot(str, result=str)
    def createFolder(self, name):
        try:
            ident = self.store.create_folder(name)
            self.foldersChanged.emit()
            self.notice.emit("أُنشئ المجلد «" + name.strip() + "».", False)
            return ident
        except Exception as exc:
            self.notice.emit(str(exc), True)
            return ""

    @Slot(str, str)
    def assignJobFolder(self, job_id, folder_id):
        try:
            self.store.assign_folder(job_id, folder_id)
            self.refresh()
            def work():
                try:
                    self.scheduler.library.publish_current(self.store, job_id, self.config.snapshot())
                    self.notice.emit("حُدّث مجلد المهمة ونسختها في المكتبة.", False)
                except Exception as exc:
                    self.notice.emit("حُفظ تصنيف المهمة، وتعذر تحديث نسخة المكتبة: " + str(exc), True)
                finally:
                    self.refreshLibrary()
            threading.Thread(target=work, daemon=True, name="folder-assignment").start()
        except Exception as exc:
            self.notice.emit(str(exc), True)

    @Property(bool, notify=libraryChanged)
    def libraryBusy(self): return self._library_busy

    @Property(bool, notify=youtubeVersionChanged)
    def youtubeUpdating(self): return self._youtube_updating

    @Property(str, notify=youtubeVersionChanged)
    def youtubeVersion(self):
        from .youtube_update import binary_version
        try:
            active = binary_version(self.config.home)
            if active:
                return active + " (محدّثة)"
        except Exception:
            pass
        import yt_dlp
        return yt_dlp.version.__version__ + " (مدمجة)"

    @Slot()
    def updateYoutube(self):
        if self._youtube_updating:
            return
        if not self.scheduler.idle() or self._busy:
            self.notice.emit("انتظر اكتمال قراءة المصدر أو تنزيله قبل تحديث أداة YouTube.", True)
            return
        self._youtube_updating = True
        self.youtubeVersionChanged.emit()
        def work():
            try:
                from .youtube_update import update
                with self.scheduler.youtube.gate:
                    version = update(self.config.home)
                self.notice.emit("اكتمل تحديث yt-dlp إلى " + version + " وسيُستخدم في الطلبات التالية.", False)
            except Exception as exc:
                self.notice.emit("فشل تحديث yt-dlp: " + str(exc), True)
            finally:
                self._youtube_updating = False
                self.youtubeVersionChanged.emit()
        threading.Thread(target=work, daemon=True, name="youtube-update").start()

    @Property("QVariantMap", notify=changed)
    def stats(self):
        return {"total": len(self._jobs), "active": sum(j["status"] == "running" for j in self._jobs),
                "done": sum(j["status"] == "completed" for j in self._jobs),
                "attention": sum(j["status"] in ("failed", "blocked", "uncertain") for j in self._jobs),
                "documents": sum(bool(j["document"]) for j in self._jobs)}

    @Slot()
    def refresh(self):
        self._publish_accounts()
        result = []
        for j in self.store.jobs():
            stages = j["stages"]
            if all(s["status"] in ("completed", "skipped") for s in stages.values()):
                state, current = "completed", "opencode"
            else:
                current = next(s for s in STAGES if stages[s]["status"] not in ("completed", "skipped"))
                state = stages[current]["status"]
            result.append({k: v for k, v in j.items() if k not in ("config", "stages") } | {
                "status": state, "statusLabel": LABELS.get(state, state), "current": current,
                "stageLabel": STAGE_LABELS[current], "error": stages[current]["error"],
                "worker": stages[current]["worker"], "progress": stages[current]["progress"],
                "retryAt": self.retry_at(stages[current]["next_retry"])
                           if state == "pending" and stages[current]["next_retry"] > time.time() else "",
                "template": j["config"].get("template_name", ""),
                "steps": [{"name": name, "label": STAGE_LABELS[name], "status": stages[name]["status"]}
                          for name in (("opencode",) if str(j.get("input_path", "")).lower().endswith(".txt") else STAGES)],
            })
        workers = self.scheduler.workers()
        for w in workers:
            w["stageLabel"] = STAGE_LABELS[w["stage"]]
        logs = [{**r, "time": time.strftime("%H:%M:%S", time.localtime(r["ts"]))} for r in self.store.events(100)]
        running = self.scheduler.running
        if self._auto_stop and not self._auto_stop_fired:
            if self.store.queue_drained():
                # Fire only after real work was seen, so arming on an empty queue never quits immediately.
                if (self._auto_stop_seen and running and not self._busy and not self.remote.busy
                        and not self.internet_commands.busy and not self.internet.busy and not self.accountBusy and self.scheduler.idle()):
                    self._auto_stop_fired = True
                    self.store.event("", "اكتمل تنفيذ الطابور؛ سيُغلق التطبيق تلقائيًا.", "info")
                    self.autoStopReached.emit()
            else:
                self._auto_stop_seen = True
        # The timer fires while idle too; republishing identical state would
        # make QML re-evaluate every binding for nothing.
        jobs_changed = result != self._jobs
        if not jobs_changed and workers == self._workers and logs == self._logs and running == self._running:
            return
        self._jobs, self._workers, self._logs, self._running = result, workers, logs, running
        if jobs_changed:
            self._jobs_json = json.dumps(result, ensure_ascii=False)
            self._jobs_model.update_items(result)
        self._workers_model.update_items(workers)
        self._publish_results()
        if jobs_changed:
            self.jobsChanged.emit()
        self.changed.emit()

    @staticmethod
    def retry_at(timestamp):
        fmt = "%H:%M" if time.localtime(timestamp)[:3] == time.localtime()[:3] else "%Y-%m-%d %H:%M"
        return time.strftime(fmt, time.localtime(timestamp))

    @Slot(str, result="QVariantMap")
    def jobById(self, ident):
        return next((j for j in self._jobs if j["id"] == ident), {})

    def _publish_results(self):
        query = self._results_query.lower()
        rows = [j for j in self._jobs
                if (j["document"] or j["transcript"] or j["audio"])
                and (not query or query in str(j["title"]).lower())
                and (self._results_folder == "all" or j.get("folder_id", "") == self._results_folder)]
        self._results_model.update_items(rows)

    @Slot(str)
    def setResultsQuery(self, text):
        text = str(text)
        if text == self._results_query:
            return
        self._results_query = text
        self._publish_results()

    @Slot(str)
    def setResultsFolder(self, ident):
        self._results_folder = ident
        self._publish_results()

    def _publish_files(self):
        query = self._files_query.lower()
        rows = [f for f in self._library
                if (self._files_kind == "all" or f["kind"] == self._files_kind)
                and (not query or query in str(f["name"]).lower())
                and (self._files_folder == "all" or f["group"] == self._files_folder)]
        self._files_model.update_items(rows)

    @Slot(str)
    def setFilesQuery(self, text):
        text = str(text)
        if text == self._files_query:
            return
        self._files_query = text
        self._publish_files()

    @Slot(str)
    def setFilesKind(self, kind):
        kind = str(kind)
        if kind == self._files_kind:
            return
        self._files_kind = kind
        self._publish_files()

    @Slot(str)
    def setFilesFolder(self, name):
        self._files_folder = name
        self._publish_files()

    def _provider_rows(self):
        """Built-in providers plus the ones the user added; only the added ones are editable."""
        builtin = {item["id"] for item in BUILTIN_PROVIDERS}
        return [{"id": p["id"], "name": p["name"], "keyEnv": p["key_env"], "isolated": p["isolated"],
                 "builtin": p["id"] in builtin} for p in provider_catalog(self.config.snapshot())]

    @Slot()
    def _publish_providers(self):
        self._providers_model.update_items(self._provider_rows())

    @Slot(str)
    def _accept_source(self, raw):
        data = json.loads(raw)
        self._source = data["items"]
        self._source_title = data["title"]
        self._busy = False
        self._source_model.update_items(self._source)
        self.sourceChanged.emit()

    @Slot(str)
    def _error(self, message):
        self._busy = False
        self.sourceChanged.emit()
        self.notice.emit(message, True)

    @Slot(str, bool)
    def discover(self, url, refresh):
        if self._busy:
            return
        self._busy = True
        self.sourceChanged.emit()
        self.discovery_cancel = threading.Event()
        cfg = self.config.snapshot()
        def work():
            try:
                data = self.scheduler.youtube.discover(url, cfg, self.discovery_cancel, refresh)
                self.sourceReady.emit(json.dumps(data, ensure_ascii=False))
            except Exception as exc:
                self.backgroundError.emit(str(exc))
        threading.Thread(target=work, daemon=True).start()

    @Slot()
    def cancelDiscovery(self):
        self.discovery_cancel.set()

    @Slot(int, bool)
    def selectSource(self, index, selected):
        if 0 <= index < len(self._source):
            self._source[index]["selected"] = selected and self._source[index]["available"]
            self._source_model.update_items(self._source)
            self.sourceChanged.emit()

    @Slot(str)
    def selectRange(self, spec):
        try:
            wanted = set()
            if spec.strip() not in ("all", "none"):
                for part in spec.replace("،", ",").split(","):
                    bits = [int(s.strip()) for s in part.split("-")]
                    if len(bits) == 1:
                        wanted.add(bits[0])
                    elif len(bits) == 2 and 0 < bits[0] <= bits[1] <= 100000:
                        wanted.update(range(bits[0], bits[1] + 1))
                    else:
                        raise ValueError()
            for row in self._source:
                row["selected"] = row["available"] and (spec == "all" or row["position"] in wanted)
            self._source_model.update_items(self._source)
            self.sourceChanged.emit()
        except ValueError:
            self.notice.emit("اكتب نطاقًا مثل 1-10,15,20-25", True)

    @Slot(int, int)
    def moveSource(self, index, direction):
        target = index + direction
        if 0 <= target < len(self._source):
            self._source[index], self._source[target] = self._source[target], self._source[index]
            # Same length, so this lands as a dataChanged on two rows: the rows
            # swap content in place and the view neither jumps nor flickers.
            self._source_model.update_items(self._source)
            self.sourceChanged.emit()

    @Slot(str, str)
    @Slot(str, str, str)
    def enqueue(self, preset, template_id="", folder_id=""):
        try:
            cfg = self.config.snapshot()
            if preset != "custom":
                cfg["download_enabled"] = True
                cfg["deepgram_enabled"] = preset != "download"
                cfg["opencode_enabled"] = preset == "full"
            template = next((t for t in cfg["opencode_templates"] if t["id"] == template_id), None)
            if template:
                cfg["prompt"] = template["text"]
                cfg["template_name"] = template["name"]
            selected = [dict(row) for row in self._source if row["selected"]]
            if not selected:
                raise ValueError("حدد فيديو واحدًا على الأقل.")
            # Preserve original lesson numbers; Store saves the visual queue order separately.
            added = self.store.add(selected, cfg, folder_id=folder_id)
            self.notice.emit(f"أضيفت {len(added)} مهمة. اضغط تشغيل لبدء المعالجة.", False)
            self.refresh()
        except Exception as exc:
            self.notice.emit(str(exc), True)

    @Slot(result=int)
    @Slot(str, result=int)
    @Slot(str, str, result=int)
    def importLocal(self, folder_id="", template_id=""):
        paths, _ = QFileDialog.getOpenFileNames(None, "استيراد صوت أو تفريغ", "", "Media (*.mp3 *.wav *.m4a *.flac *.ogg *.txt)")
        return self._import_paths(paths, template_id, folder_id, transcripts_only=False)

    @Slot(result=int)
    @Slot(str, result=int)
    @Slot(str, str, result=int)
    def importTranscripts(self, template_id="", folder_id=""):
        paths, _ = QFileDialog.getOpenFileNames(None, "إضافة تفريغات إلى الطابور", "", "تفريغات TXT (*.txt *.TXT)")
        return self.importTranscriptPaths(paths, template_id, folder_id)

    @staticmethod
    def _transcript_order(path):
        # Tagged segments make names beginning with digits and letters comparable.
        return tuple((1, int(part)) if part.isdigit() else (0, part.casefold())
                     for part in re.split(r"(\d+)", path.name))

    @Slot(result=int)
    @Slot(str, result=int)
    @Slot(str, str, result=int)
    def importTranscriptFolder(self, template_id="", folder_id=""):
        name = QFileDialog.getExistingDirectory(None, "اختيار مجلد التفريغات", "")
        if not name:
            return 0
        try:
            paths = sorted((path for path in Path(name).iterdir() if path.is_file() and path.suffix.lower() == ".txt"),
                           key=lambda path: (self._transcript_order(path), path.name.casefold(), path.name))
        except OSError as exc:
            self.notice.emit("تعذر قراءة مجلد التفريغات: " + str(exc), True)
            return 0
        if not paths:
            self.notice.emit("لا توجد ملفات TXT مباشرة في المجلد المحدد.", True)
            return 0
        return self.importTranscriptPaths(paths, template_id, folder_id)

    @Slot("QVariantList", result=int)
    @Slot("QVariantList", str, result=int)
    @Slot("QVariantList", str, str, result=int)
    def importDroppedTranscripts(self, urls, template_id="", folder_id=""):
        paths, errors = [], []
        for value in urls:
            url = value if isinstance(value, QUrl) else QUrl(str(value))
            if not url.isLocalFile() or url.hasQuery() or url.hasFragment():
                errors.append("يُقبل سحب ملفات TXT المحلية فقط.")
                continue
            paths.append(url.toLocalFile())
        return self._import_paths(paths, template_id, folder_id, errors=errors)

    def importTranscriptPaths(self, paths, template_id="", folder_id=""):
        return self._import_paths(paths, template_id, folder_id)

    def _import_paths(self, paths, template_id, folder_id, transcripts_only=True, errors=None):
        errors = list(errors or [])
        items, seen = [], set()
        for name in paths:
            try:
                path = Path(name).resolve(strict=True)
                canonical = os.path.normcase(str(path))
                if canonical in seen:
                    continue
                seen.add(canonical)
                if not path.is_file():
                    raise ValueError("المسار ليس ملفًا.")
                extensions = (".txt",) if transcripts_only else (".txt", ".mp3", ".wav", ".m4a", ".flac", ".ogg")
                if path.suffix.lower() not in extensions:
                    raise ValueError("صيغة الملف غير مدعومة؛ أضف ملف TXT." if transcripts_only else "صيغة الملف غير مدعومة.")
                if path.suffix.lower() == ".txt":
                    if not path.read_text(encoding="utf-8-sig").strip():
                        raise ValueError("ملف التفريغ فارغ.")
                items.append({"video_id": "local-" + path.stem, "title": path.stem, "position": len(items) + 1,
                              "url": "", "input_path": str(path)})
            except (OSError, UnicodeError, ValueError) as exc:
                errors.append(f"تعذر استيراد «{Path(name).name}»: {exc}")
        count = 0
        if items:
            try:
                cfg = self.config.snapshot()
                cfg["download_enabled"] = False
                cfg["deepgram_enabled"] = any(not item["input_path"].lower().endswith(".txt") for item in items)
                cfg["opencode_enabled"] = True
                template = next((t for t in cfg["opencode_templates"] if t["id"] == template_id), None)
                if template_id and not template:
                    raise ValueError("القالب المحدد غير موجود؛ اختر قالبًا آخر.")
                if template:
                    cfg["prompt"], cfg["template_name"] = template["text"], template["name"]
                count = len(self.store.add(items, cfg, folder_id=folder_id))
                self.refresh()
            except Exception as exc:
                errors.append("تعذر إضافة الملفات إلى الطابور: " + str(exc))
                errors.extend(getattr(exc, "__notes__", ()))
        if count or errors:
            message = f"أضيفت {count} ملفات إلى الطابور." if count else "لم تُضف ملفات إلى الطابور."
            if errors:
                message += "\n" + "\n".join(dict.fromkeys(errors))
            self.notice.emit(message, bool(errors))
        return count

    @Slot()
    def toggleRunning(self):
        if self.scheduler.running:
            self.scheduler.pause()
            self.notice.emit("توقف توزيع المهام. العمال النشطون يكملون مهامهم الحالية.", False)
        else:
            self.scheduler.start()
        self.refresh()

    @Slot(bool)
    def setAutoStop(self, enabled):
        """Session-only switch: when the queue fully drains, the app quits itself."""
        self._auto_stop = bool(enabled)
        if not enabled:
            self._auto_stop_seen = self._auto_stop_fired = False

    @Slot(str, bool)
    def pauseStage(self, name, paused):
        self.scheduler.pause_stage(name, paused)
        self.refresh()

    @Slot(str)
    def recheckStage(self, stage):
        """Re-check a service after its automatic block was fixed: lift the block
        and requeue the jobs the block had parked."""
        self.scheduler.unblock_stage(stage)
        for job in self.store.jobs():
            record = job["stages"].get(stage)
            if record and record["status"] == "blocked":
                self.store.set_stage(job["id"], stage, "pending")
                self.store.event(job["id"], "أُعيد فحص الخدمة فأُعيدت المهمة إلى الطابور.")
        self.notice.emit("أُعيد فحص الخدمة؛ ستُتابع المهام العمل تلقائيًا إذا كان الطابور يعمل.", False)
        self.refresh()

    @Slot(str)
    def cancelJob(self, ident):
        self.scheduler.cancel_job(ident)
        self.refresh()

    @Slot("QVariantList", result=int)
    def deleteJobs(self, idents):
        try:
            result = self.scheduler.delete_jobs(idents)
            count = len(result["deleted"])
            self.refresh()
            message = f"حُذف {count} عنصر من الطابور. الملفات والمخرجات محفوظة."
            if result["busy"]:
                message += f" تعذر حذف {len(result['busy'])} عنصر نشط؛ ألغِه وانتظر توقفه ثم أعد الحذف."
            self.notice.emit(message, bool(result["busy"]))
            return count
        except Exception as exc:
            self.notice.emit("تعذر حذف عناصر الطابور: " + str(exc), True)
            return 0

    @Slot(str, str, bool)
    def retry(self, ident, stage, rerun):
        try:
            self._retry_one(ident, stage, rerun)
            # Retrying a download lifts the automatic block (cookies, disk) but
            # never the user's manual pause and never the 429 cooldown.
            manual = stage == "download" and stage in self.scheduler.manual_paused_stages
            if stage == "download":
                self.scheduler.unblock_stage(stage)
            remaining = max(0, int(self.scheduler.youtube.cooldown_until - time.time()))
            if stage == "download" and remaining > 0:
                self.notice.emit(f"أُعيدت المهمة إلى الطابور وستبدأ تلقائيًا بعد انتهاء تهدئة YouTube ({self.scheduler.human_delay(remaining)}).", False)
            elif manual:
                self.notice.emit("أُعيدت المهمة إلى الطابور، لكن خدمة التنزيل موقوفة يدويًا من صفحة العمال.", False)
            elif not self.scheduler.running:
                self.notice.emit("أُعيدت المهمة إلى الطابور. شغّل الطابور للبدء.", False)
            else:
                self.notice.emit("أُعيدت المهمة إلى الطابور وسيبدأ المنظم بتنفيذها.", False)
            self.refresh()
        except Exception as exc:
            self.notice.emit(str(exc), True)

    def _retry_one(self, ident, stage, rerun=False):
        if stage == "opencode" and ident in self.store.state_get("chatgpt_session_holds", {}):
            raise ValueError("تحقّق من توقف جلسة ChatGPT السابقة عبر «فحص الجلسة السابقة» في صفحة العمال، ثم أعد المحاولة.")
        self.scheduler.retry_job(ident, stage, rerun)
        if not rerun:
            job = self.store.get(ident)
            for name in STAGES[STAGES.index(stage) + 1:]:
                if job["stages"][name]["status"] == "cancelled":
                    self.store.set_stage(ident, name, "pending")

    @Slot('QVariantList', str, str)
    def retryMany(self, idents, stage, status):
        """Apply one action to matching attention jobs; ignore stale UI selections."""
        if stage not in STAGES or status not in ("failed", "blocked", "uncertain"):
            self.notice.emit("اختر مجموعة مهام متشابهة لإعادة المحاولة.", True)
            return
        count = skipped = 0
        for ident in dict.fromkeys(str(value) for value in idents):
            job = self.store.get(ident)
            if not job or any(record["status"] == "running" for record in job["stages"].values()):
                skipped += 1
                continue
            current = next((name for name in STAGES if job["stages"][name]["status"] not in ("completed", "skipped")), "")
            if current != stage or job["stages"][stage]["status"] != status:
                skipped += 1
                continue
            try:
                self._retry_one(ident, stage)
                self.store.event(ident, "إعادة محاولة جماعية بطلب المستخدم.")
                count += 1
            except Exception:
                skipped += 1
        if count and stage == "download":
            self.scheduler.unblock_stage(stage)
        reason = "تحتاج فحص الجلسة السابقة أو تغيّرت حالتها" if stage == "opencode" and status == "uncertain" else "تغيّرت حالتها"
        self.notice.emit(f"أُعيدت {count} مهمة إلى الطابور" + (f"؛ تخطيت {skipped} مهمة {reason}." if skipped else "."), not bool(count))
        self.refresh()

    @Slot(str)
    def openPath(self, path):
        if path and Path(path).exists():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(path).resolve())))
        else:
            self.notice.emit("الملف أو المجلد غير موجود بعد.", True)

    @Slot(str)
    def copyText(self, text):
        QGuiApplication.clipboard().setText(str(text))
        self.notice.emit("نُسخ المسار إلى الحافظة.", False)

    @Slot(str, result=str)
    def choosePath(self, kind):
        if kind == "folder":
            return QFileDialog.getExistingDirectory(None, "اختيار مجلد")
        path, _ = QFileDialog.getOpenFileName(None, "اختيار ملف", "", "All files (*)")
        return path

    @Slot(str, result=bool)
    def saveSettings(self, raw):
        try:
            self.config.save(json.loads(raw))
            self.scheduler.disabled_workers.clear()
            self.scheduler.forget_keys()
            self.settingsChanged.emit()
            self.notice.emit("حُفظت الإعدادات. المهام الجديدة تستخدم إعدادات المعالجة الجديدة.", False)
            return True
        except Exception as exc:
            self.notice.emit(str(exc), True)
            return False

    @Slot(str, str)
    def saveKey(self, worker, key):
        try:
            stage, provider = self._worker_scope(worker)
            self.secrets.set(credential_name(stage, worker, provider), key)
            self.scheduler.forget_keys()
            self.scheduler.disabled_workers.pop(worker, None)
            self.notice.emit("حُفظ المفتاح في مخزن بيانات اعتماد Windows.", False)
            self.refresh()
        except Exception as exc:
            self.notice.emit("تعذر حفظ المفتاح: " + str(exc), True)

    def _worker_scope(self, worker):
        """Stage and configured provider of a worker id, used for the credential name."""
        data = self.config.snapshot()
        for item in data["deepgram_workers"]:
            if item["id"] == worker:
                return "deepgram", ""
        for item in data["opencode_workers"]:
            if item["id"] == worker:
                return "opencode", provider_of(item, data)
        raise ValueError("لم يُعثر على العامل.")

    @Slot(str, bool)
    def setWorkerEnabled(self, worker, enabled):
        """Start or stop this worker taking new jobs, straight from the worker list."""
        data = self.config.snapshot()
        for kind in ("deepgram_workers", "opencode_workers"):
            for item in data[kind]:
                if item["id"] == worker:
                    item["enabled"] = bool(enabled)
                    self.config.save(data)
                    self.scheduler.disabled_workers.clear()
                    self.settingsChanged.emit()
                    self.notice.emit(("سُجّل العامل " if enabled else "أُوقف العامل ") + item["name"] +
                                     (" لاستقبال المهام." if enabled else " عن استقبال مهام جديدة."), False)
                    self.refresh()
                    return
        self.notice.emit("لم يُعثر على العامل.", True)

    @Slot(str, str, str)
    def setProviderOverride(self, provider, model, variant):
        """Session-only: run every OpenCode worker on one provider, model and reasoning level."""
        data = self.config.snapshot()
        self.scheduler.set_provider_override(provider, model, variant)
        if provider:
            note = f"سيُفرض مزود {provider_info(provider, data)['name']} على عمال OpenCode."
            if variant:
                note += f" مستوى التفكير: {variant}."
            missing = [w["name"] for w in data["opencode_workers"]
                       if key_env_for(provider, data) and not self.scheduler.key_state("opencode", w, provider)]
            if missing:
                note += " بلا مفتاح حتى الآن: " + "، ".join(missing) + "."
            self.notice.emit(note, False)
        else:
            self.notice.emit("أُلغي الفرض؛ كل عامل يعود إلى إعداده الخاص.", False)
        self.refresh()

    @Slot(str, result=bool)
    def saveProvider(self, raw):
        """Add or update a user-defined provider; built-ins are never touched."""
        try:
            item = json.loads(raw)
            data = self.config.snapshot()
            entry = {"id": str(item.get("id", "")).strip(), "name": str(item.get("name", "")).strip() or str(item.get("id", "")).strip(),
                     "key_env": str(item.get("keyEnv", "")).strip(), "isolated": bool(item.get("isolated", True))}
            data["opencode_providers"] = [p for p in data["opencode_providers"] if p["id"] != entry["id"]] + [entry]
            self.config.save(data)
            self.scheduler.forget_keys()
            self.settingsChanged.emit()
            self.notice.emit("حُفظ المزود " + entry["name"] + ".", False)
            self.refresh()
            return True
        except Exception as exc:
            self.notice.emit(str(exc), True)
            return False

    @Slot(str)
    def deleteProvider(self, provider):
        try:
            data = self.config.snapshot()
            used = [w["name"] for w in data["opencode_workers"] if provider_of(w, data) == provider]
            if used:
                self.notice.emit("لا يمكن حذف مزود تستخدمه: " + "، ".join(used), True)
                return
            data["opencode_providers"] = [p for p in data["opencode_providers"] if p["id"] != provider]
            self.config.save(data)
            self.scheduler.forget_keys()
            self.settingsChanged.emit()
            self.notice.emit("حُذف المزود.", False)
            self.refresh()
        except Exception as exc:
            self.notice.emit(str(exc), True)

    @Slot(str)
    def resetWorker(self, worker):
        self.scheduler.disabled_workers.pop(worker, None)
        self.scheduler.forget_keys()
        self.refresh()

    @Slot()
    def refreshLibrary(self):
        """Rescan the library folders. Called when the file page opens, not on the UI timer."""
        try:
            self._library = self.scheduler.library.scan(self.config.snapshot())
        except Exception as exc:
            self._library = []
            self.notice.emit("تعذر قراءة المكتبة: " + str(exc), True)
        self.libraryChanged.emit()

    @Slot()
    def organizeLibrary(self):
        """Copy every finished artifact into the library, off the UI thread."""
        if self._library_busy:
            return
        self._library_busy = True
        self.libraryChanged.emit()
        def work():
            try:
                copied = self.scheduler.library.organize_current(self.store, self.config.snapshot())
                self.notice.emit(f"نُظّم {copied} ملفًا في المكتبة.", False)
            except Exception as exc:
                self.notice.emit("تعذر التنظيم: " + str(exc), True)
            finally:
                self._library_busy = False
                self.refreshLibrary()
        threading.Thread(target=work, daemon=True, name="library").start()

    @Slot()
    def exportWord(self):
        destination = QFileDialog.getExistingDirectory(None, "تصدير ملفات Word إلى مجلد", "")
        if not destination or self._library_busy:
            return
        self._library_busy = True
        self.libraryChanged.emit()
        def work():
            try:
                count = self.scheduler.library.export_documents(self.store.jobs(), self.config.snapshot(), destination)
                self.notice.emit(f"صُدّر {count} مستند Word إلى {destination} وفق مجلدات المواد والدروس.", False)
            except Exception as exc:
                self.notice.emit("تعذر تصدير Word: " + str(exc), True)
            finally:
                self._library_busy = False
                self.libraryChanged.emit()
        threading.Thread(target=work, daemon=True, name="word-export").start()

    @Slot(str)
    def deleteLibraryFile(self, path):
        try:
            self.scheduler.library.remove(self.config.snapshot(), path)
            self.notice.emit("حُذف الملف من المكتبة؛ النسخة الأصلية في مجلد المهمة باقية.", False)
        except Exception as exc:
            self.notice.emit(str(exc), True)
        self.refreshLibrary()

    @Slot()
    def importLegacy(self):
        path, _ = QFileDialog.getOpenFileName(None, "استيراد إعدادات السكربت", str(Path.home()), "JSON (*.json)")
        if path:
            try:
                self.config.import_legacy(Path(path), self.secrets)
                self.scheduler.disabled_workers.clear()
                self.scheduler.forget_keys()
                self.settingsChanged.emit()
                self.notice.emit("استوردت الإعدادات ومفتاح العامل الأول. أضف مفتاحًا مستقلًا للعامل الثاني.", False)
            except Exception as exc:
                self.notice.emit(str(exc), True)

    @Slot()
    def checkEnvironment(self):
        def work():
            import importlib.util
            import shutil
            messages = []
            for name in ("yt_dlp", "yt_dlp_ejs", "imageio_ffmpeg", "requests", "keyring"):
                messages.append(name + (": جاهز" if importlib.util.find_spec(name) else ": غير مثبت"))
            messages.append("JavaScript: " + (shutil.which("node") or shutil.which("deno") or "غير موجود"))
            try:
                executable = resolve_opencode(self.config.data["opencode_executable"])
                messages.append("OpenCode " + opencode_version(executable) + ": " + str(executable))
            except Exception as exc:
                messages.append(str(exc))
            self.backgroundError.emit("\n".join(messages))
        threading.Thread(target=work, daemon=True).start()

    @Slot(result=bool)
    def prepareClose(self):
        from PySide6.QtWidgets import QMessageBox
        if not self.scheduler.idle() or self._busy or self.remote.busy or self.internet_commands.busy or self.internet.busy or self.accountBusy:
            answer = QMessageBox.question(None, "إغلاق ورق", "توجد عمليات نشطة. الإغلاق يوقف العمليات المحلية وقد يترك طلب Deepgram بنتيجة غير مؤكدة. هل تريد الإغلاق؟")
            if answer != QMessageBox.StandardButton.Yes:
                return False
        self.discovery_cancel.set()
        self._update_cancel.set()
        self.remote.stop()
        self.stopInternet()
        self.closeRemoteLibrary()
        self.closeAccounts()
        self.scheduler.shutdown()
        return True

    @Slot()
    def exportReport(self):
        path, _ = QFileDialog.getSaveFileName(None, "حفظ تقرير", "تقرير-المعالجة.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()
            ws = wb.active
            ws.title = "المعالجة"
            ws.sheet_view.rightToLeft = True
            ws.append(["الترتيب", "العنوان", "الحالة", "المرحلة", "الرابط", "Word", "الخطأ"])
            for j in self._jobs:
                # Do not interpret untrusted titles as spreadsheet formulas.
                values = [j["position"], j["title"], j["statusLabel"], j["stageLabel"], j["url"], j["document"], j["error"]]
                ws.append(values)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="142338")
            for col in "ABCDEFG":
                ws.column_dimensions[col].width = 35 if col != "B" else 65
            ws.freeze_panes = "A2"
            wb.save(path)
            self.notice.emit("حُفظ تقرير المعالجة.", False)
        except Exception as exc:
            self.notice.emit(str(exc), True)
